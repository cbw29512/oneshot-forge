import docx, copy, re, sys
SRC="repo/assets/bundles/owlbear-arms/07_Source/Right_To_Bear_Arms_Editable_DM_Packet.docx"
OUT="oa12/Right_To_Bear_Arms_Editable_DM_Packet_v1.2.docx"
d=docx.Document(SRC)
P=d.paragraphs; T=d.tables
log=[]
def setp(p,text):
    """replace paragraph text, keeping first run formatting"""
    runs=p.runs
    if not runs: p.add_run(text); return
    runs[0].text=text
    for r in runs[1:]: r.text=""
def para(i,old_sub,text):
    p=P[i]; assert old_sub in p.text, (i,p.text[:80]); setp(p,text); log.append(f"P{i}")
def cellp(ti,r,c,j,old_sub,text):
    p=T[ti].cell(r,c).paragraphs[j]; assert old_sub in p.text,(ti,r,c,j,p.text[:80]); setp(p,text); log.append(f"T{ti}r{r}c{c}p{j}")
def drop_table(ti):
    el=T[ti]._tbl; el.getparent().remove(el); log.append(f"-T{ti}")
def drop_para(i):
    el=P[i]._p; el.getparent().remove(el); log.append(f"-P{i}")

# --- title / version
para(2,"A One-Shot for 5 Players","A One-Shot for 3–6 Players (5 heroes provided), Level 6 · v1.2")
cellp(0,0,1,0,"5 players","3–6 players (tuned for 5; see Part 1, §6 scaling)")
# --- Mantle name (v1.2 §8: Magistrate's Mantle, CUSTOM)
for p in P:
    if "Mantle of Intellect" in p.text: setp(p,p.text.replace("Mantle of Intellect","Magistrate's Mantle")); log.append("mantle")
for t in T:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                if "Mantle of Intellect" in p.text: setp(p,p.text.replace("Mantle of Intellect","Magistrate's Mantle"))
# --- chronology & Edra motive (v1.2 §1)
para(39,"Three months ago","Three months ago, Edra sold the raiders one patrol rotation — the night the north logging camp was supposed to be empty. It wasn't. Her younger brother Tomas was working there. He survived the raid, saw signs that someone inside Brindlebarrow had given the raiders the patrol routes, and started asking questions. Edra fed him false leads to keep him off her trail. By then the raiders owned her, and she has kept feeding them ever since.")
para(40,"She does not know","Two days ago, Tomas and Old Cobb followed a fresh raider trail into the Western Woods. Near the edge of Grubbleclaw's lair, Tomas found physical evidence tying Edra to the raiders. Edra's raider contact intercepted him after he and Cobb separated, and killed him there. Tomas died in the woods carrying the proof — not in the logging-camp raid, and not at Grubbleclaw's claws. Edra knows only that he has not come back. She refuses to believe he is dead.")
cellp(5,0,0,0,"","Why this lands hard: Edra is not a monster. She is a competent leader who made one terrible choice and has spent three months covering it — even from her own brother. The party will trust her. When they figure it out, the question is not 'do we kill the traitor' but 'what do you do with someone whose cover-up got her brother killed.' That is the second moral question of the session, layered under the first.") if len(T[5].cell(0,0).paragraphs)==1 else None
# T5 may be multi-paragraph: rewrite all paragraphs
c=T[5].cell(0,0)
texts=["Why this lands hard","Edra is not a monster. She is a competent leader who made one terrible choice and has spent three months covering it — even from her own brother. The party will trust her. When they figure it out, the question is not 'do we kill the traitor' but 'what do you do with someone whose cover-up got her brother killed.' That is the second moral question of the session, layered under the first."]
ps=c.paragraphs
for k,p in enumerate(ps):
    setp(p, texts[k] if k<len(texts) else "")
log.append("T5")
# Cobb (T12)
cellp(12,3,0,4,"If thanked for the map","If thanked for the map. \"Don't thank me. I left a man in there. Tomas Vell. We split up at the edge of the bird's ground and when I found him again he was gone. I could not carry him. The eggs were chirping. I ran. ... Tell Edra I'm sorry, would you?\"")
cellp(12,3,0,5,"DM Note","DM Note. Cobb is the one who can break the case wide open if the party asks the right question. He followed a raider trail into the woods with Tomas — Edra's brother — two days ago. He believes the owlbear killed Tomas because that is where he last saw him alive, but if pressed about the body he admits: \"No claw marks. Blade wounds. And his coat was cut open, like someone wanted what he was carrying.\" Tomas died in the woods at raider hands, carrying evidence against Edra. Edra does not know this yet.")
# T12 role line
for j,p in enumerate(T[12].cell(1,0).paragraphs):
    if "Survivor of a previous attempt" in p.text:
        setp(p,p.text.replace("Survivor of a previous attempt on the lair.","Went into the woods with Tomas Vell two days ago and came back alone.")); log.append("T12role")
# Edra (T13)
cellp(13,1,0,1,"Wants","Wants  Her secret kept — and her brother Tomas, missing two days, to be alive. (He is dead. She has not let herself believe it.)")
cellp(13,3,0,0,"Stat block reference","Stat block reference. Use the SRD Veteran (CR 3, 700 XP): AC 17, HP 58, Multiattack two longsword attacks, Longsword +5 (1d8+3), Heavy Crossbow +3 (1d10+1).")
cellp(13,3,0,5,"If asked about her scar","If asked about her scar. (Touches it without thinking.) \"Last raid. Three months back. I was lucky. My brother was at the camp too — he got out. I don't talk about it.\"")
cellp(13,3,0,6,"If asked about Tomas","If asked about Tomas (her brother). \"He went into the woods with Cobb two days ago. He's alive. I know he is. Don't — don't ask me how I know. He's alive.\"")
cellp(13,3,0,7,"When confronted with evidence","When confronted with evidence (private). (Long silence. Then she sits down. She does not draw her sword.) \"...how long have you known? Was it the boots? The map? Both?\" (A pause.) \"One patrol rotation. That's all it was, the first time. The camp was supposed to be empty. Tomas was there. He lived, and he started asking who knew the rotation. I sent him north, south, anywhere but at me. Then he went into the woods with Cobb.\" (Quietly.) \"He's alive. Tell me he's alive.\"")
cellp(13,3,0,8,"When told Tomas is dead","When told Tomas is dead (by Cobb or evidence). (She does not speak. She removes her cloak, folds it, lays it on the ground. Picks up her sword. Looks at the party.) \"Then they killed him for what he found on me. I have nothing left to lose. Tell me where to die. I will die useful.\"")
# Skitters (T20/T21)
cellp(21,0,0,4,"TRAITOR CLUE","• TRAITOR CLUE — If asked how the raiders know so much about Brindlebarrow's defenses: 'The iron-haired one with the scar. She told us where to hit. South gate. Every week she leaves the patrol times under a stone on the south road. I fetch some of them.'")
cellp(21,0,0,5,"If pressed further","• If pressed further on the iron-haired woman: 'Don't know her name. Tall. Has a limp she hides. Brings information, takes coin back. Always at the south road, never the north. Two days ago the boss sent people after some man in the woods who found her notes. They came back with his coat.'")
# Grubbleclaw combat (P169 + T35)
para(169,"Grubbleclaw is scaled","v1.2: run Grubbleclaw from the combat block in Part 1, §8 (AC 16 with the Magistrate's Mantle, 15 without · HP 68 · Speed 40 ft. · about CR 4, 1,100 XP; an upgraded SRD Owlbear). He will not fight to the death unless protecting the eggs. Removing the mantle requires Grubbleclaw to be restrained or grappled, then a full action and a successful DC 16 Strength (Athletics) or Dexterity (Sleight of Hand) check; failure alerts and enrages him. Removal changes only his AC.")
drop_table(35)
# Edra escapes / mid-battle timing
para(182,"If she escapes","If she escapes: she reaches the raider lines. She returns in Wave 3 at the village, fighting beside the Raid Captain, in place of one Ash Berserker. Use the SRD Veteran.")
para(187,"Round 3 of the Front A battle","Wave 2 at the village (Part 1, §3): Edra is holding the south gate. The party will see her step BACK from the line at a critical moment, opening a hole. Insight DC 13 to recognize the hesitation as deliberate (not panic).")
para(188,"Round 4","Next round: a raider yells past her, 'WHERE IS OUR INSIDE WOMAN?' Edra does not respond. The implication will land.")
# Cobb reveal (T37) and confession (T36)
cellp(36,0,0,6,"He sent me a letter","\"One patrol rotation. That's all it was, the first time. The camp was supposed to be empty. Tomas was there. He lived — and he started asking who knew the rotation. I sent him the wrong way, again and again. Then he went into the woods with Cobb.\"")
cellp(36,0,0,8,"He IS alive","\"...He's alive. He has to be. Tell me he's alive.\"")
cellp(37,0,0,2,"Tomas didn't make it out","\"Edra. Edra, I have to. Tomas is dead. Two days ago, in the woods, at the edge of the owlbear's ground. We split up. When I found him it wasn't claws. It was blades. His coat was cut open, like someone wanted what he was carrying. I left him because I could not carry him. I told myself it was the owlbear. I told myself a lot of things.\"")
cellp(37,0,0,8,"corpse's ring","\"...three months. I sold one night, and it cost me my brother.\"")
# The Two Fronts / Egg-Raid Decision / Village DP -> v1.2 procedure
para(197,"The Two Fronts","The Final Defense — run it from Part 1")
para(198,"Front A","Use Part 1, §2–6 for every number")
para(199,"Main assault","Four Preparation Hours; Brindlebarrow starts at 4 Defense Points (minimum 0, maximum 6); Routes A–E decide who defends the village and who defends the nest; each front runs three waves of three rounds; the wave rosters, XP math, and 3–6 player scaling are in §6; the quick-run enemy blocks (Raider, Raider Scout, Ash Berserker, Ash Shaman, Raid Captain) are in §7.")
para(200,"Front B","Story beats to keep")
para(201,"Triggers at the start","On Wave 3, the egg team breaks for the nest while everyone is engaged at the gate — Skitters (if alive) or a village kid shouts that figures are breaking west. Moving the eggs into the village chapel needs Grubbleclaw's consent (Route D). The best ending is earned when the party and Grubbleclaw defend Brindlebarrow together.")
for ti in (43,42,41,40,39): drop_table(ti)
for i in range(213,201,-1): drop_para(i)
# Endings / aftermath
para(246,"Bonus reward","Bonus reward: Cobb gives the party Tomas's pinky ring to lay on Edra's grave beside her. He weeps for the first time since the woods.")
para(248,"Campaign hook","Campaign hook. She is alive somewhere, neither raider nor villager. The raiders' Ash Shaman will hunt her. So will Brindlebarrow.")
para(252,"Hand this to the party","Hand this to the party when they meet Old Cobb in Brindlebarrow — an aging hunter who followed a raider trail into the woods with Tomas Vell two days ago, was webbed at the cave mouth, knocked out, and dragged half-conscious to where he eventually crawled out. His memory of the lair is partial and the map gets less reliable the deeper it goes.")
para(269,"Tomas Vell","Tomas Vell: survived the logging-camp raid three months ago; killed in the woods two days ago by Edra's raider contact while carrying evidence of her betrayal. Grubbleclaw did not kill him.")
# cheat sheet
para(278,"Grubbleclaw at a glance","Grubbleclaw at a glance — Part 1, §8: AC 16 (15 without mantle) · HP 68 · Beak +8 (1d10+5) · Claw +8 (2d8+5) · Javelin +8 (1d6+5)")
para(279,"Village Defense quick math","Village Defense quick math — Part 1, §2 and §5: start 4 DP · min 0 · max 6 · Bess drills +1 · fortify +1")
drop_table(51); drop_table(50)
para(290,"He sent me a letter","'One patrol rotation. That's all it was, the first time.'")
para(291,"corpse's ring","'I sold one night, and it cost me my brother.'")
# T4 clocks row
cellp(4,3,2,0,"Defense Points start low","Fewer Defense Points (Part 1, §2)")
d.save(OUT); print(len(log),"edits; saved",OUT)
# --- second pass fixes
d=docx.Document(OUT)
def allps(d):
    for p in d.paragraphs: yield p
    for t in d.tables:
        seen=set()
        for r in t.rows:
            for c in r.cells:
                if id(c._tc) in seen: continue
                seen.add(id(c._tc))
                for p in c.paragraphs: yield p
n=0
for p in allps(d):
    s=p.text
    if "I can teach a turnip" in s:
        setp(p,"If asked to train villagers. \"I can teach a turnip to swing a spear in two hours. Get me a turnip and two hours. Or six villagers. Same thing.\" (Bess drills the militia: 2 Preparation Hours, +1 Defense Point — Part 1, §2.)"); n+=1
    elif "drop from 3 to 1" in s:
        setp(p,s.replace("Their starting Village Defense Points drop from 3 to 1 (see Final Battle).","Brindlebarrow starts the defense at 2 Defense Points instead of 4 (Part 1, §2).")); n+=1
d.save(OUT); print("second pass",n)
d=docx.Document(OUT)
for s in d.sections:
    for p in s.header.paragraphs:
        if "DM Packet" in p.text: setp(p,"Right to OwlBear Arms v1.2 — Part 2: Scenes, NPCs & Dialogue")
    for p in s.footer.paragraphs:
        if p.text.strip()=="Page" and p.runs: p.runs[0].text="Part 2 · page "
# intro note under the subtitle
P2=d.paragraphs
for p in P2:
    if p.text.startswith("A One-Shot for 3–6 Players"):
        note=copy.deepcopy(p._p); p._p.addnext(note)
        from docx.text.paragraph import Paragraph
        np_=Paragraph(note,p._parent); setp(np_,"Part 2 of the v1.2 DM Book. Rules, numbers, and the chronology are in Part 1 — if anything here seems to differ, Part 1 wins.")
        break
d.save(OUT); print("headers ok")
d=docx.Document(OUT)
for p in d.paragraphs:
    if p.text.strip()=="RIGHT TO BEAR ARMS": setp(p,"RIGHT TO OWLBEAR ARMS"); break
d.save(OUT)
